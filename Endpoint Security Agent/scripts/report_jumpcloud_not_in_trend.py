"""
Cron script: dispositivos en JumpCloud que NO están en Trend Vision One.
Compara por hostname normalizado. Guarda Excel en reports/ y actualiza .meta.json.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from agent.core.config import get_settings
from agent.core.logging import configure_logging
from agent.clients.trend_vision_one import TrendVisionOneClient
from agent.clients.jumpcloud import JumpCloudClient
from agent.services.slack import notify

REPORTS_DIR = ROOT / "reports"
META_FILE   = REPORTS_DIR / ".meta.json"
REPORT_KEY  = "jumpcloud_not_in_trend"


def normalize(name: str) -> str:
    return name.strip().lower().split(".")[0]


def fetch_trend_hostnames(settings) -> set[str]:
    with TrendVisionOneClient(settings) as client:
        items = list(client.fetch_collection(
            "/v3.0/endpointSecurity/endpoints",
            {"top": "100"},
            max_pages=25,
        ))
    return {normalize(item.get("endpointName", "")) for item in items if item.get("endpointName")}


def get_field(fields: list, label: str):
    for f in fields:
        if f.get("label") == label:
            dt = f.get("dataType")
            if dt == "text":
                return f.get("text", {}).get("value", "")
            if dt == "reference":
                return (f.get("reference", {}).get("value") or {}).get("name", "")
            if dt == "boolean":
                return f.get("boolean", {}).get("value")
    return None


def fetch_jumpcloud_laptops(settings) -> list[dict]:
    with JumpCloudClient(settings) as jc:
        devices = list(jc.fetch_collection("/v2/asset-management/devices", limit=100, max_pages=50))

    laptops = []
    for d in devices:
        fields = d.get("fields", [])
        if (
            get_field(fields, "Type") == "Laptop"
            and get_field(fields, "Status") == "In Use"
            and get_field(fields, "Owner")
        ):
            name = get_field(fields, "Name") or ""
            os_raw = get_field(fields, "OS Family") or get_field(fields, "Operating System (OS)") or ""
            laptops.append({
                "hostname":     name,
                "hostname_key": normalize(name),
                "os":           "macOS" if os_raw.lower() in ("darwin", "mac") else os_raw,
                "model":        get_field(fields, "Model") or "—",
                "serial":       get_field(fields, "Serial Number") or "—",
                "os_version":   get_field(fields, "OS Version") or "—",
                "owner":        get_field(fields, "Owner") or "—",
                "last_contact": (get_field(fields, "Last Contact") or "")[:19].replace("T", " ") or "—",
                "agent_status": get_field(fields, "Agent Status") or "—",
                "mdm_status":   get_field(fields, "MDM Status") or "—",
                "disk_enc":     get_field(fields, "Disk Encrypted"),
            })
    return laptops


def fetch_data(settings) -> list[dict]:
    print("Consultando Trend Vision One...")
    trend_hostnames = fetch_trend_hostnames(settings)
    print(f"  Trend: {len(trend_hostnames)} endpoints")

    print("Consultando JumpCloud...")
    jc_laptops = fetch_jumpcloud_laptops(settings)
    print(f"  JumpCloud (Laptops In Use): {len(jc_laptops)} dispositivos")

    missing = [d for d in jc_laptops if d["hostname_key"] not in trend_hostnames]
    missing.sort(key=lambda x: x["hostname"].lower())
    return missing


def build_excel(rows: list[dict], output_path: Path) -> None:
    YELLOW = "FDFA3D"; BLACK = "0A0A0A"; GRAY = "F2F2F2"
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = ["#", "Hostname", "Sistema Operativo", "Modelo", "Serial",
               "OS Version", "Owner", "Último Contacto", "Agent Status", "MDM Status", "Disco Enc."]

    wb = Workbook()
    ws = wb.active
    ws.title = "JumpCloud sin Trend"

    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")

    ws.merge_cells("A1:K1")
    t = ws["A1"]
    t.value = "JumpCloud — Laptops In Use no encontradas en Trend Vision One"
    t.font = Font(name="Arial", bold=True, size=14, color=BLACK)
    t.fill = PatternFill("solid", start_color=YELLOW)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:K2")
    s = ws["A2"]
    s.value = f"Generado: {now_str}  |  Datos en tiempo real  |  Total: {len(rows)} equipos"
    s.font = Font(name="Arial", size=9, color="666666")
    s.fill = PatternFill("solid", start_color="FAFAFA")
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=BLACK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[3].height = 22

    for i, row in enumerate(rows, 1):
        r = 3 + i
        bg = GRAY if i % 2 == 0 else "FFFFFF"
        enc = row["disk_enc"]
        enc_display = "Sí" if enc is True else ("No" if enc is False else "—")
        vals = [i, row["hostname"], row["os"], row["model"], row["serial"],
                row["os_version"], row["owner"], row["last_contact"],
                row["agent_status"], row["mdm_status"], enc_display]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name="Arial", size=9, color=BLACK)
            cell.fill = PatternFill("solid", start_color=bg)
            cell.alignment = Alignment(vertical="center",
                                       horizontal="center" if col in (1, 3, 9, 10, 11) else "left")
            cell.border = border
            if col == 11 and enc_display == "No":
                cell.font = Font(name="Arial", size=9, bold=True, color="CC0000")
        ws.row_dimensions[r].height = 16

    tr = 3 + len(rows) + 1
    ws.merge_cells(f"A{tr}:B{tr}")
    ws[f"A{tr}"].value = f"Total: {len(rows)} equipos"
    ws[f"A{tr}"].font = Font(name="Arial", bold=True, size=10, color=BLACK)
    ws[f"A{tr}"].fill = PatternFill("solid", start_color=YELLOW)
    ws[f"A{tr}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[tr].height = 20

    win = sum(1 for r in rows if "windows" in r["os"].lower())
    mac = len(rows) - win
    ws.merge_cells(f"C{tr}:F{tr}")
    ws[f"C{tr}"].value = f"Windows: {win}   |   macOS: {mac}"
    ws[f"C{tr}"].font = Font(name="Arial", bold=True, size=10, color=BLACK)
    ws[f"C{tr}"].fill = PatternFill("solid", start_color=YELLOW)
    ws[f"C{tr}"].alignment = Alignment(horizontal="center", vertical="center")

    for col, w in enumerate([4, 28, 14, 28, 20, 13, 28, 20, 14, 13, 12], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A4"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def update_meta(row_count: int, output_path: Path) -> None:
    meta = {}
    if META_FILE.exists():
        try:
            meta = json.loads(META_FILE.read_text())
        except Exception:
            pass
    meta[REPORT_KEY] = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "row_count":    row_count,
        "filename":     output_path.name,
    }
    META_FILE.write_text(json.dumps(meta, indent=2))


def main() -> None:
    configure_logging()
    settings = get_settings()

    rows = fetch_data(settings)
    print(f"Encontrados: {len(rows)} laptops en JumpCloud que NO están en Trend")

    ts = datetime.now().strftime("%Y-%m-%d")
    output_path = REPORTS_DIR / f"{REPORT_KEY}_{ts}.xlsx"
    build_excel(rows, output_path)
    update_meta(len(rows), output_path)
    print(f"Reporte guardado: {output_path}")

    win = sum(1 for r in rows if "windows" in r["os"].lower())
    mac = len(rows) - win
    notify(
        report_key=REPORT_KEY,
        row_count=len(rows),
        output_path=output_path,
        extra_lines=[f"🖥️  Windows: {win}   |   🍎 macOS: {mac}"],
    )


if __name__ == "__main__":
    main()
