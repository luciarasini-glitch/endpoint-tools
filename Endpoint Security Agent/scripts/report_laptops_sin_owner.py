"""
Cron script: genera el reporte de laptops In Use sin owner asignado.
Guarda el Excel en reports/ y actualiza reports/.meta.json.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from agent.core.config import get_settings
from agent.core.logging import configure_logging
from agent.clients.jumpcloud import JumpCloudClient
from agent.services.slack import notify

REPORTS_DIR = ROOT / "reports"
META_FILE   = REPORTS_DIR / ".meta.json"
REPORT_KEY  = "laptops_in_use_sin_owner"


def get_field(fields: list, label: str):
    for f in fields:
        if f.get("label") == label:
            dt = f.get("dataType")
            if dt == "text":         return f.get("text", {}).get("value", "")
            if dt == "reference":    return (f.get("reference", {}).get("value") or {}).get("name", "")
            if dt == "boolean":      return f.get("boolean", {}).get("value")
            if dt == "referenceList":
                return [v.get("name", "") for v in (f.get("referenceList", {}).get("value") or [])]
    return None


def fetch_data(settings) -> list[dict]:
    rows = []
    with JumpCloudClient(settings) as jc:
        all_devices = list(jc.fetch_collection("/v2/asset-management/devices", limit=100, max_pages=50))

    for d in all_devices:
        fields = d.get("fields", [])
        if (
            get_field(fields, "Type") == "Laptop"
            and get_field(fields, "Status") == "In Use"
            and not get_field(fields, "Owner")
        ):
            os_raw = get_field(fields, "OS Family") or get_field(fields, "Operating System (OS)") or ""
            rows.append({
                "Nombre":          get_field(fields, "Name") or "",
                "Sistema Operativo": "macOS" if os_raw.lower() in ("darwin", "mac") else os_raw,
                "Modelo":          get_field(fields, "Model") or "",
                "Vendor":          get_field(fields, "Vendor") or "",
                "Serial Number":   get_field(fields, "Serial Number") or "",
                "OS Version":      get_field(fields, "OS Version") or "",
                "Ultimo Contacto": (get_field(fields, "Last Contact") or "")[:19].replace("T", " ") or "—",
                "Agent Status":    get_field(fields, "Agent Status") or "",
                "MDM Status":      get_field(fields, "MDM Status") or "",
                "Disk Encrypted":  get_field(fields, "Disk Encrypted"),
            })
    return rows


def build_excel(rows: list[dict], output_path: Path) -> None:
    YELLOW = "FDFA3D"
    BLACK  = "0A0A0A"
    GRAY   = "F2F2F2"
    thin   = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["#", "Nombre", "Sistema Operativo", "Modelo", "Vendor",
               "Serial Number", "OS Version", "Ultimo Contacto",
               "Agent Status", "MDM Status", "Disco Encriptado"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Laptops Sin Asignar"

    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")

    ws.merge_cells("A1:K1")
    t = ws["A1"]
    t.value = "Laptops In Use — Sin Owner Asignado"
    t.font = Font(name="Arial", bold=True, size=14, color=BLACK)
    t.fill = PatternFill("solid", start_color=YELLOW)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:K2")
    s = ws["A2"]
    s.value = f"Generado: {now_str}  |  Total: {len(rows)} equipos  |  Filtros: Type=Laptop, Status=In Use, Owner=vacío"
    s.font = Font(name="Arial", size=9, color="666666")
    s.fill = PatternFill("solid", start_color="FAFAFA")
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=BLACK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[3].height = 22

    for i, row in enumerate(rows, 1):
        r = 3 + i
        bg = GRAY if i % 2 == 0 else "FFFFFF"
        enc = row.get("Disk Encrypted")
        enc_display = "Sí" if enc is True else ("No" if enc is False else "—")
        vals = [i, row["Nombre"], row["Sistema Operativo"], row["Modelo"],
                row["Vendor"], row["Serial Number"], row["OS Version"],
                row["Ultimo Contacto"], row["Agent Status"], row["MDM Status"], enc_display]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name="Arial", size=9, color=BLACK)
            cell.fill = PatternFill("solid", start_color=bg)
            cell.alignment = Alignment(vertical="center",
                                       horizontal="center" if col in (1, 3, 5, 9, 10, 11) else "left")
            cell.border = border
            if col == 11 and enc_display == "No":
                cell.font = Font(name="Arial", size=9, bold=True, color="CC0000")
        ws.row_dimensions[r].height = 16

    tr = 3 + len(rows) + 1
    ws.merge_cells(f"A{tr}:B{tr}")
    ws[f"A{tr}"].value = f"Total: {len(rows)} equipos"
    ws[f"A{tr}"].font = Font(name="Arial", bold=True, size=10, color=BLACK)
    ws[f"A{tr}"].fill = PatternFill("solid", start_color=YELLOW)
    ws[f"A{tr}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[tr].height = 20

    win = sum(1 for r in rows if "windows" in r["Sistema Operativo"].lower())
    mac = len(rows) - win
    ws.merge_cells(f"C{tr}:F{tr}")
    ws[f"C{tr}"].value = f"Windows: {win}   |   macOS: {mac}"
    ws[f"C{tr}"].font = Font(name="Arial", bold=True, size=10, color=BLACK)
    ws[f"C{tr}"].fill = PatternFill("solid", start_color=YELLOW)
    ws[f"C{tr}"].alignment = Alignment(horizontal="center", vertical="center")

    for col, w in enumerate([4, 36, 14, 32, 10, 24, 12, 22, 13, 13, 17], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A4"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def update_meta(row_count: int, output_path: Path) -> None:
    meta = {}
    if META_FILE.exists():
        try:
            meta = json.loads(META_FILE.read_text())
        except Exception:
            pass
    meta[REPORT_KEY] = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "row_count":    row_count,
        "filename":     output_path.name,
    }
    META_FILE.write_text(json.dumps(meta, indent=2))


def main() -> None:
    configure_logging()
    settings = get_settings()

    print("Consultando JumpCloud...")
    rows = fetch_data(settings)
    print(f"Encontrados: {len(rows)} laptops In Use sin owner")

    ts = datetime.now().strftime("%Y-%m-%d")
    output_path = REPORTS_DIR / f"{REPORT_KEY}_{ts}.xlsx"
    build_excel(rows, output_path)
    update_meta(len(rows), output_path)
    print(f"Reporte guardado: {output_path}")

    win = sum(1 for r in rows if "windows" in r["Sistema Operativo"].lower())
    mac = len(rows) - win
    notify(
        report_key=REPORT_KEY,
        row_count=len(rows),
        output_path=output_path,
        extra_lines=[f"🖥️  Windows: {win}   |   🍎 macOS: {mac}"],
    )


if __name__ == "__main__":
    main()
