"""
Cron script: genera el reporte de endpoints Trend Vision One sin conexión hace +10 días.
Guarda el Excel en reports/ y actualiza reports/.meta.json.
"""
from __future__ import annotations

import json
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(ROOT))

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from agent.core.config import get_settings
from agent.core.logging import configure_logging
from agent.clients.trend_vision_one import TrendVisionOneClient
from agent.services.slack import notify

REPORTS_DIR = ROOT / "reports"
META_FILE   = REPORTS_DIR / ".meta.json"
REPORT_KEY  = "trend_endpoints_sin_conexion_10d"


def parse_dt(val: str | None) -> datetime | None:
    if not val:
        return None
    try:
        dt = datetime.fromisoformat(val.replace("Z", "+00:00"))
        return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
    except Exception:
        return None


def fetch_data(settings) -> list[dict]:
    cutoff = datetime.now(timezone.utc) - timedelta(days=10)
    stale = []

    with TrendVisionOneClient(settings) as client:
        all_items = list(client.fetch_collection(
            "/v3.0/endpointSecurity/endpoints",
            {"top": "100"},
            max_pages=25,
        ))

    for item in all_items:
        edr = item.get("edrSensor") or {}
        epp = item.get("eppAgent") or {}
        last_edr = parse_dt(edr.get("lastConnectedDateTime"))
        last_epp = parse_dt(epp.get("lastConnectedDateTime"))
        last = max(filter(None, [last_edr, last_epp]), default=None)

        if last and last < cutoff:
            days_ago = (datetime.now(timezone.utc) - last).days
            stale.append({
                "hostname":     item.get("endpointName", ""),
                "os":           item.get("osName", ""),
                "platform":     item.get("osPlatform", "").capitalize(),
                "ip":           (item.get("ipAddresses") or ["—"])[0],
                "last_seen":    last.strftime("%d/%m/%Y %H:%M"),
                "days_offline": days_ago,
                "last_user":    item.get("lastLoggedOnUser", "") or "—",
                "edr_status":   edr.get("connectivity", "—"),
                "epp_status":   epp.get("status", "—"),
                "serial":       item.get("serialNumber", "") or "—",
            })

    stale.sort(key=lambda x: x["days_offline"], reverse=True)
    return stale


def build_excel(rows: list[dict], output_path: Path) -> None:
    YELLOW = "FDFA3D"; BLACK = "0A0A0A"; GRAY = "F2F2F2"
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    headers = ["#", "Hostname", "OS", "Plataforma", "IP", "Última Conexión",
               "Días Offline", "Último Usuario", "EDR Status", "EPP Status", "Serial"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Sin Conexion +10 dias"

    now_str = datetime.now().strftime("%d/%m/%Y %H:%M")

    ws.merge_cells("A1:K1")
    t = ws["A1"]
    t.value = "Trend Vision One — Endpoints sin conexión hace más de 10 días"
    t.font = Font(name="Arial", bold=True, size=14, color=BLACK)
    t.fill = PatternFill("solid", start_color=YELLOW)
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    ws.merge_cells("A2:K2")
    s = ws["A2"]
    s.value = f"Generado: {now_str}  |  Datos en tiempo real desde Trend Vision One  |  Total: {len(rows)} endpoints"
    s.font = Font(name="Arial", size=9, color="666666")
    s.fill = PatternFill("solid", start_color="FAFAFA")
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 18

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=3, column=col, value=h)
        cell.font = Font(name="Arial", bold=True, size=10, color="FFFFFF")
        cell.fill = PatternFill("solid", start_color=BLACK)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = border
    ws.row_dimensions[3].height = 22

    for i, row in enumerate(rows, 1):
        r = 3 + i
        bg = GRAY if i % 2 == 0 else "FFFFFF"
        vals = [i, row["hostname"], row["os"], row["platform"], row["ip"],
                row["last_seen"], row["days_offline"], row["last_user"],
                row["edr_status"], row["epp_status"], row["serial"]]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name="Arial", size=9, color=BLACK)
            cell.fill = PatternFill("solid", start_color=bg)
            cell.alignment = Alignment(vertical="center",
                                       horizontal="center" if col in (1, 4, 7, 9, 10) else "left")
            cell.border = border
            if col == 7 and isinstance(val, int) and val > 30:
                cell.font = Font(name="Arial", size=9, bold=True, color="CC0000")
        ws.row_dimensions[r].height = 16

    tr = 3 + len(rows) + 1
    ws.merge_cells(f"A{tr}:B{tr}")
    ws[f"A{tr}"].value = f"Total: {len(rows)} endpoints"
    ws[f"A{tr}"].font = Font(name="Arial", bold=True, size=10, color=BLACK)
    ws[f"A{tr}"].fill = PatternFill("solid", start_color=YELLOW)
    ws[f"A{tr}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[tr].height = 20

    win = sum(1 for r in rows if "windows" in r["platform"].lower())
    mac = len(rows) - win
    ws.merge_cells(f"C{tr}:F{tr}")
    ws[f"C{tr}"].value = f"Windows: {win}   |   macOS: {mac}"
    ws[f"C{tr}"].font = Font(name="Arial", bold=True, size=10, color=BLACK)
    ws[f"C{tr}"].fill = PatternFill("solid", start_color=YELLOW)
    ws[f"C{tr}"].alignment = Alignment(horizontal="center", vertical="center")

    for col, w in enumerate([4, 28, 16, 12, 16, 20, 13, 30, 13, 12, 18], 1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = "A4"
    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def update_meta(row_count: int, output_path: Path) -> None:
    meta = {}
    if META_FILE.exists():
        try:
            meta = json.loads(META_FILE.read_text())
        except Exception:
            pass
    meta[REPORT_KEY] = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "row_count":    row_count,
        "filename":     output_path.name,
    }
    META_FILE.write_text(json.dumps(meta, indent=2))


def main() -> None:
    configure_logging()
    settings = get_settings()

    print("Consultando Trend Vision One...")
    rows = fetch_data(settings)
    print(f"Encontrados: {len(rows)} endpoints sin conexión hace +10 días")

    ts = datetime.now().strftime("%Y-%m-%d")
    output_path = REPORTS_DIR / f"{REPORT_KEY}_{ts}.xlsx"
    build_excel(rows, output_path)
    update_meta(len(rows), output_path)
    print(f"Reporte guardado: {output_path}")

    win = sum(1 for r in rows if "windows" in r["platform"].lower())
    mac = len(rows) - win
    notify(
        report_key=REPORT_KEY,
        row_count=len(rows),
        output_path=output_path,
        extra_lines=[f"🖥️  Windows: {win}   |   🍎 macOS: {mac}"],
    )


if __name__ == "__main__":
    main()
